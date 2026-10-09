import os
import re
import html
import contextvars

import pymupdf as fitz
import docx
import openpyxl
from pptx import Presentation
from sqlalchemy import text

from config import PROJECT_ROOT, FILESYSTEM_CONFIG, ACTIVE_TABLES
from database import db_engine

_table_text_cols = {}
_current_role = contextvars.ContextVar("current_role", default="guest")


def exec_query_table_data(table_name: str, keyword: str = "", role: str = "guest",
                          columns: str = "*", category: str = "",
                          order_by: str = "", limit: int = 8) -> dict:
    if not db_engine:
        return {"error": "Database offline."}
    if role != "admin" and ACTIVE_TABLES and table_name not in ACTIVE_TABLES:
        return {"error": f"Table '{table_name}' not in whitelist."}
    if role != "admin" and table_name == "cerita":
        return {"error": "Access denied. Admin only."}

    limit = max(1, min(int(limit), 20))
    allowed_cols = ACTIVE_TABLES.get(table_name, [])

    try:
        with db_engine.connect() as conn:
            if columns.strip() == "*" or not columns.strip():
                select_part = "*"
            else:
                requested = [c.strip().lower() for c in columns.split(",") if c.strip()]
                safe_cols = requested if role == "admin" else [
                    c for c in requested if not allowed_cols or c in allowed_cols
                ]
                if not safe_cols:
                    return {"error": "No valid columns requested."}
                select_part = ", ".join(f"`{c}`" for c in safe_cols)

            order_part = ""
            if order_by.strip():
                parts = order_by.strip().split()
                col_name = parts[0].lower()
                direction = parts[1].upper() if len(parts) > 1 else "ASC"
                if direction not in ("ASC", "DESC"):
                    direction = "ASC"
                if role == "admin" or not allowed_cols or col_name in allowed_cols:
                    order_part = f" ORDER BY `{col_name}` {direction}"

            where_clauses = []
            params = {}

            if category.strip():
                where_clauses.append("`kategori` = :category")
                params["category"] = category.strip()

            if keyword.strip():
                if table_name not in _table_text_cols:
                    res = conn.execute(text(f"SHOW COLUMNS FROM `{table_name}`"))
                    _table_text_cols[table_name] = [
                        r[0] for r in res.fetchall()
                        if "char" in r[1].lower() or "text" in r[1].lower()
                    ]
                text_cols = _table_text_cols[table_name]
                if text_cols:
                    for i, kw in enumerate(k.strip() for k in keyword.split(",") if k.strip()):
                        col_conds = [f"`{c}` LIKE :kw_{i}" for c in text_cols]
                        where_clauses.append("(" + " OR ".join(col_conds) + ")")
                        params[f"kw_{i}"] = f"%{kw}%"

            where = (" WHERE " + " AND ".join(where_clauses)) if where_clauses else ""
            sql = f"SELECT {select_part} FROM `{table_name}`{where}{order_part} LIMIT {limit}"

            data_res = conn.execute(text(sql), params)
            keys = list(data_res.keys())
            rows = data_res.fetchall()

            if not rows:
                return {"note": f"No matching data in '{table_name}'."}

            csv_lines = ["|".join(keys)]
            for r in rows:
                cleaned_row = []
                for idx, val in enumerate(r):
                    text_val = str(val) if val is not None else ""

                    if role != "admin":
                        col_lower = keys[idx].lower()
                        if ("name" in col_lower and "app" not in col_lower) or col_lower == "bank_name":
                            text_val = "xxx"
                        elif col_lower in ("content", "review_text", "konten"):
                            text_val = re.sub(
                                r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', 'xxx', text_val
                            )
                            text_val = re.sub(r'\b(?:\d{1,3}\.){3}\d{1,3}\b', 'xxx', text_val)
                            text_val = re.sub(r'\b(?:\+62|62|0)[2-9][0-9]{7,11}\b', 'xxx', text_val)

                    text_val = html.unescape(text_val)
                    text_val = re.sub(r'<[^>]+>', ' ', text_val)
                    text_val = re.sub(r'\s+', ' ', text_val).strip()
                    cleaned_row.append(text_val)
                csv_lines.append("|".join(cleaned_row))

            return {"data": "\n".join(csv_lines)}
    except Exception as e:
        return {"error": str(e)}


def list_files(subfolder: str = "") -> dict:
    base_dir = os.path.join(PROJECT_ROOT, FILESYSTEM_CONFIG["doc_folder"])
    target_dir = os.path.join(base_dir, subfolder).strip()
    if not os.path.exists(target_dir):
        return {"error": f"Folder not found: {target_dir}"}
    try:
        return {"files": [
            f for f in os.listdir(target_dir)
            if os.path.isfile(os.path.join(target_dir, f))
            and any(f.lower().endswith(ext) for ext in FILESYSTEM_CONFIG["allowed_ext"])
        ]}
    except Exception as e:
        return {"error": str(e)}


def read_file(filename: str, subfolder: str = "") -> dict:
    base_dir = os.path.join(PROJECT_ROOT, FILESYSTEM_CONFIG["doc_folder"])
    filepath = os.path.join(base_dir, subfolder, filename).strip()
    if not os.path.exists(filepath):
        return {"error": f"File not found: {filename}"}

    ext = os.path.splitext(filename)[1].lower()
    if ext not in FILESYSTEM_CONFIG["allowed_ext"]:
        return {"error": f"Extension '{ext}' not allowed."}

    try:
        content = ""
        if ext in (".txt", ".csv"):
            with open(filepath, "r", encoding="utf-8") as f:
                content = f.read()
        elif ext == ".pdf":
            doc = fitz.open(filepath)
            for page in doc:
                content += page.get_text()
            doc.close()
        elif ext == ".docx":
            doc = docx.Document(filepath)
            content = "\n".join(p.text for p in doc.paragraphs)
        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(filepath, data_only=True)
            for sheet in wb.sheetnames:
                content += f"\n--- Sheet: {sheet} ---\n"
                ws = wb[sheet]
                for row in ws.iter_rows(values_only=True):
                    content += " | ".join(str(c) if c is not None else "" for c in row) + "\n"
        elif ext == ".pptx":
            prs = Presentation(filepath)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        content += shape.text + "\n"

        if len(content) > 15000:
            content = content[:15000] + "...(truncated)"
        return {"content": content.strip() or "File is empty."}
    except Exception as e:
        return {"error": f"Failed to read file: {e}"}


def query_table_data(table_name: str, keyword: str = "", columns: str = "*",
                     category: str = "", order_by: str = "", limit: int = 8) -> dict:
    return exec_query_table_data(
        table_name, keyword, _current_role.get(), columns, category, order_by, limit
    )
